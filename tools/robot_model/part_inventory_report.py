#!/usr/bin/env python3
"""Create the review table for every printed or prototype-aluminium CAD body.

The current CAD is one-sided for the leg and is mirrored by the robot builder.  Therefore
this inventory keeps one row per CAD body and states the physical multiplicity separately.
The mass column is a *body-level input* to URDF link aggregation, not necessarily a URDF
link of the same name.
"""
from __future__ import annotations

import hashlib
import html
import json
import os
from pathlib import Path
import shutil

import numpy as np
from PIL import Image, ImageDraw, ImageFont
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from alu_parts_render import render
from massprops_fusion import classify


REPO = Path(__file__).resolve().parents[2]
SNAP = REPO / "tools/robot_model/fusion_snapshots/v30_inspection"
OUT = SNAP / "part_inventory"
IMG = OUT / "img"
AL_DENSITY = 2.70
CURRENT_MAP = {
    "AnkleA2B_A_L", "AnkleA2B_A_L (1)", "AnkleA2B_A_R",
    "AnkleA2B_A_R (1)", "AnkleA2B_A_Support", "AnkleA2B_A_Support (1)",
    "Arm_A", "Arm_B", "FEET",
}


def short_name(path: str) -> str:
    return path.rsplit("::", 1)[-1]


def occurrence(path: str) -> str:
    return path.rsplit("/", 1)[-1].split("::", 1)[0]


def quantity(path: str) -> tuple[int, str]:
    if "/Joints_UnderBody:1/CenterParts:1" in path:
        return 1, "center body"
    if "/Joints_UnderBody:1/" in path:
        return 2, "one-leg CAD mirrored L/R"
    if "/Arm_R:1/" in path:
        return 2, "right-arm CAD mirrored L/R"
    return 1, "single CAD placement"


def density_rgb(density: float) -> str:
    # low-density print = warm yellow/green, 6061 = cool blue; clamp exotic values.
    t = max(0.0, min(1.0, (density - 0.75) / (AL_DENSITY - 0.75)))
    a, b = np.array([245, 224, 132]), np.array([142, 190, 232])
    c = np.rint(a * (1 - t) + b * t).astype(int)
    return "#%02x%02x%02x" % tuple(c)


def legacy_images() -> list[dict]:
    rows = []
    for root in (Path("/home/syaro/pyg_fea/fusion/alu_parts_v5"),
                 Path("/home/syaro/pyg_fea/fusion/alu_parts")):
        index = root / "index.json"
        if index.exists():
            for row in json.loads(index.read_text()):
                row = dict(row)
                row["_root"] = root
                rows.append(row)
    return rows


def find_legacy(path: str, rows: list[dict]) -> Path | None:
    name, occ = short_name(path), occurrence(path).split(":")[0]
    hits = [r for r in rows if r.get("body") == name]
    if len(hits) > 1:
        scoped = [r for r in hits if r.get("occ", "").split(":")[0] == occ]
        hits = scoped or hits
    for row in hits:
        p = Path(row.get("img", ""))
        if p.exists():
            return p
        p = row["_root"] / "img" / f'{row["key"]}.png'
        if p.exists():
            return p
    return None


def make_image(path: str, out: Path, legacy: list[dict]) -> str:
    key = hashlib.sha256(path.encode()).hexdigest()[:20]
    mesh = SNAP / "meshes/.parts" / f"{key}.npz"
    if mesh.exists():
        z = np.load(mesh)
        v, f = np.asarray(z["v"], float).reshape(-1, 3), np.asarray(z["f"], int).reshape(-1, 3)
        render(v - v.mean(0), f, str(out), 520)
        return "current v30 cached mesh"
    old = find_legacy(path, legacy)
    if old:
        shutil.copy2(old, out)
        return f"legacy exact-name render ({old.parent.parent.name})"
    canvas = Image.new("RGB", (520, 390), "white")
    ImageDraw.Draw(canvas).text((30, 175), "MESH RENDER MISSING", fill="red")
    canvas.save(out)
    return "render missing"


def label_card(row: dict, width: int = 430, height: int = 340) -> Image.Image:
    card = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(card)
    draw.rectangle((0, 0, width - 1, height - 1), outline="#8a96a3", width=2)
    src = Image.open(row["image_abs"]).convert("RGB")
    src.thumbnail((width - 20, height - 82))
    card.paste(src, ((width - src.width) // 2, 42 + (height - 82 - src.height) // 2))
    draw.rectangle((0, 0, width, 42), fill=row["density_color"])
    draw.text((10, 7), row["name"], fill="black")
    draw.text((10, height - 30),
              f'{row["volume_cm3"]:.2f} cm3 | now {row["mass_g"]:.2f} g | AL {row["al_mass_g"]:.2f} g',
              fill="#1c2733")
    return card


def mapping_sheet(rows: list[dict], names: set[str], title: str, note: str, path: Path) -> None:
    selected = [r for r in rows if r["name"] in names]
    cols = 3
    cw, ch, top = 430, 340, 88
    lines = (len(selected) + cols - 1) // cols
    sheet = Image.new("RGB", (cw * cols, top + ch * lines), "#edf1f4")
    draw = ImageDraw.Draw(sheet)
    draw.rectangle((0, 0, sheet.width, top), fill="#152332")
    draw.text((18, 12), title, fill="white")
    draw.text((18, 44), note, fill="#d9e4ec")
    for i, row in enumerate(selected):
        sheet.paste(label_card(row, cw, ch), ((i % cols) * cw, top + (i // cols) * ch))
    sheet.save(path)


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    IMG.mkdir(parents=True, exist_ok=True)
    printed = json.loads((SNAP / "printed_density_manifest.json").read_text())
    temporary = json.loads((SNAP / "prototype_tempmass_after_printed_manifest.json").read_text())
    original = json.loads((SNAP / "bodies.json").read_text())
    current = json.loads((SNAP / "bodies_printed_prototype_tempmass.json").read_text())
    printed_paths = {x["path"]: x for x in printed["changes"]}
    temp_paths = {x["path"]: x for x in temporary["changes"]}
    paths = list(dict.fromkeys([*printed_paths, *temp_paths]))
    legacy = legacy_images()
    rows = []
    for idx, path in enumerate(paths, 1):
        b = current[path]
        name = short_name(path)
        q, q_note = quantity(path)
        density = b["m"] * 1000 / b["v"]
        img = IMG / f"{idx:02d}_{hashlib.sha256(path.encode()).hexdigest()[:10]}.png"
        render_source = make_image(path, img, legacy)
        is_al = path in temp_paths
        state = "prototype-tempmass AL6061" if is_al else "3D-print effective density"
        mapping = "CONFIRMED-AL" if name in CURRENT_MAP else "not in confirmed AL mapping"
        note = []
        if is_al:
            note.append("worksheet AL6061 geometry mass; not a new physical weighing")
        else:
            note.append(printed_paths[path].get("ratio_source", "printed survey"))
        if name in CURRENT_MAP:
            note.append("user-confirmed blue part; one-side CAD placement mirrored L/R")
        rows.append({
            "id": idx, "name": name, "occurrence": occurrence(path), "cad_path": path,
            "owner_link": classify(path), "quantity": q, "quantity_note": q_note,
            "volume_cm3": float(b["v"]), "density_g_cm3": density,
            "density_color": density_rgb(density), "mass_g": float(b["m"] * 1000),
            "al_mass_g": float(b["v"] * AL_DENSITY),
            "original_cad_mass_g": float(original[path]["m"] * 1000),
            "material_state": state, "mapping": mapping, "notes": "; ".join(note),
            "image": str(img.relative_to(OUT)), "image_abs": str(img),
            "render_source": render_source,
        })

    # JSON is the canonical machine-readable table.
    (OUT / "part_inventory.json").write_text(
        json.dumps({
            "schema": "pygmalion-part-inventory-v1", "tag": "prototype-tempmass",
            "warning": "mass_g is a body-level input to URDF link aggregation",
            "al_density_g_cm3": AL_DENSITY, "rows": rows,
        }, ensure_ascii=False, indent=2) + "\n")

    # HTML: density cell itself carries the requested material gradient.
    body = []
    for r in rows:
        body.append(f'''<tr><td>{r["id"]}</td><td><code>{html.escape(r["name"])}</code><br><small>{html.escape(r["owner_link"])}</small></td>
<td><a href="{r["image"]}"><img src="{r["image"]}"></a></td>
<td>{r["volume_cm3"]:.4f}</td><td style="background:{r["density_color"]}">{r["density_g_cm3"]:.4f}</td>
<td><b>{r["mass_g"]:.2f}</b><br><span class="al">AL-{r["al_mass_g"]:.2f}</span></td>
<td>×{r["quantity"]}<br><small>{html.escape(r["quantity_note"])}</small></td><td>{html.escape(r["material_state"])}</td>
<td>{html.escape(r["mapping"])}</td><td>{html.escape(r["notes"])}<br><small>{html.escape(r["render_source"])}</small></td>
<td><code>{html.escape(r["cad_path"])}</code></td></tr>''')
    html_text = f'''<!doctype html><meta charset="utf-8"><title>Pygmalion part inventory</title>
<style>body{{font-family:Arial,sans-serif;color:#17222d;margin:22px}}h1{{margin-bottom:5px}}.warn{{background:#fff1c8;padding:12px;border-left:6px solid #e19b24}}table{{border-collapse:collapse;font-size:12px;width:100%}}th{{position:sticky;top:0;background:#24394b;color:white;padding:8px}}td{{border:1px solid #bbc4cc;padding:6px;vertical-align:middle}}img{{width:190px;height:142px;object-fit:contain;background:white}}code{{font-size:11px}}small{{color:#576879}}.al{{color:#1c6798}}</style>
<h1>Pygmalion printed + prototype aluminium body inventory</h1>
<p class="warn"><b>prototype-tempmass:</b> 질량은 임시 진행용입니다. “URDF 질량” 열은 독립 링크가 아니라 URDF 링크 질량으로 합산되는 CAD body 입력값입니다. 사용자가 파란 9개 CAD body를 확인했으며, 단일측 CAD를 좌우 대칭하면 18개 실물 배치입니다.</p>
<p>행 {len(rows)}개 · 단독 렌더를 누르면 원본 이미지 · 밀도색: 출력물(노랑) → Al6061(파랑)</p>
<p><a href="current_tempmass_mapping.png">확정 AL 매핑</a> · <a href="part_inventory.xlsx">Excel</a> · <a href="part_inventory.json">JSON</a></p>
<table><thead><tr><th>#</th><th>파츠명 / 집계 링크</th><th>솔로 렌더</th><th>부피<br>cm³</th><th>밀도<br>g/cm³</th><th>URDF 집계 입력<br>g / AL-g</th><th>실물 수량</th><th>재질 상태</th><th>AL 매핑</th><th>비고 / 렌더 근거</th><th>CAD body path</th></tr></thead><tbody>{''.join(body)}</tbody></table>'''
    (OUT / "index.html").write_text(html_text)

    # Spreadsheet with the same rows and embedded solo renders.
    wb = Workbook()
    ws = wb.active
    ws.title = "부품 전수표"
    headers = ["#", "파츠명", "솔로 렌더", "부피 cm³", "밀도 g/cm³", "URDF 집계 입력 g",
               "AL-6061 기준 g", "집계 링크", "실물 수량", "재질 상태", "AL 매핑", "비고", "CAD body path"]
    dark = PatternFill("solid", fgColor="24394B")
    thin = Side(style="thin", color="BBC4CC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for c, title in enumerate(headers, 1):
        x = ws.cell(1, c, title); x.fill = dark; x.font = Font(color="FFFFFF", bold=True)
        x.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for rr, r in enumerate(rows, 2):
        vals = [r["id"], r["name"], "", r["volume_cm3"], r["density_g_cm3"], r["mass_g"],
                r["al_mass_g"], r["owner_link"], r["quantity"], r["material_state"], r["mapping"],
                r["notes"] + "; " + r["render_source"], r["cad_path"]]
        for cc, val in enumerate(vals, 1):
            cell = ws.cell(rr, cc, val); cell.border = border
            cell.alignment = Alignment(vertical="center", wrap_text=cc in (2, 8, 10, 11, 12, 13))
        ws.cell(rr, 5).fill = PatternFill("solid", fgColor=r["density_color"].lstrip("#").upper())
        ws.row_dimensions[rr].height = 96
        pic = XLImage(r["image_abs"]); pic.width, pic.height = 176, 132; pic.anchor = f"C{rr}"
        ws.add_image(pic)
    for col, width in {"A":5,"B":27,"C":26,"D":13,"E":14,"F":16,"G":16,"H":20,
                       "I":11,"J":24,"K":25,"L":56,"M":68}.items():
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"; ws.auto_filter.ref = f"A1:M{len(rows)+1}"
    ws2 = wb.create_sheet("읽는 법")
    notes = [
        "ELI5: 한 CAD 바디의 부피에 현재 재질 밀도를 곱한 값이 URDF 링크 질량에 더해집니다.",
        "URDF 집계 입력 g는 독립 링크 질량이 아니라 body-level 입력입니다.",
        "AL-6061 기준 g = CAD 부피 × 2.70 g/cm³ 입니다.",
        "확정 매핑은 파란 9개 CAD body이며, 단일측 CAD를 좌우 대칭하면 18개 실물 배치입니다.",
        "FEET는 좌우 모두 알루미늄입니다. 이전 '10개' 수량은 사용자 정정으로 폐기했습니다.",
    ]
    for i, note in enumerate(notes, 1): ws2.cell(i, 1, note)
    ws2.column_dimensions["A"].width = 115
    wb.save(OUT / "part_inventory.xlsx")

    mapping_sheet(rows, CURRENT_MAP, "CONFIRMED prototype-tempmass aluminium mapping",
                  "one-side CAD: 9 blue bodies -> mirrored physical placements: 18; both FEET are aluminium",
                  OUT / "current_tempmass_mapping.png")
    print(f"{len(rows)} rows -> {OUT}")


if __name__ == "__main__":
    main()
